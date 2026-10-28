from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from clean import ORDER_KEY

CHART_COLOR = "1F4E78"
HEADER_FILL = PatternFill("solid", fgColor=CHART_COLOR)
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=12, color=CHART_COLOR)
BOLD = Font(bold=True)
TOP_LINE = Border(top=Side(style="thin"))
DATE_FORMAT = "dd-mmm-yyyy"
INT_FORMAT = "#,##0"
PCT_FORMAT = "0.0%"

SHEETS = ["Summary", "By Branch", "By Category", "Daily", "Top Products", "Clean Data", "Data Issues"]


def heading(column):
    return column.replace("_", " ").title().replace(" Id", " ID")


def money_format(currency, decimals=2):
    # quoted so Excel shows the symbol as text, whatever the system locale is
    return f'"{currency}"#,##0' + ("." + "0" * decimals if decimals else "")


def kpis(df):
    revenue = df["total"].sum()
    orders = len(df[ORDER_KEY].drop_duplicates())
    by_product = df.groupby("product")["total"].sum().sort_values(ascending=False)
    by_branch = df.groupby("branch")["total"].sum().sort_values(ascending=False)
    return {
        "revenue": round(revenue, 2),
        "orders": orders,
        "avg_order": round(revenue / orders, 2) if orders else 0,
        "units": df["quantity"].sum(),
        "top_product": (by_product.index[0], by_product.iloc[0]),
        "top_branch": (by_branch.index[0], by_branch.iloc[0]),
    }


def breakdown(df, by):
    by = [by] if isinstance(by, str) else by
    table = (
        df.assign(order=df["branch"] + "|" + df["order_id"])
        .groupby(by)
        .agg(Orders=("order", "nunique"), Units=("quantity", "sum"), Revenue=("total", "sum"))
    )
    table["Share"] = table["Revenue"] / table["Revenue"].sum()
    table = table.sort_values("Revenue", ascending=False).reset_index()
    return table.rename(columns={c: c.title() for c in by})


def daily(df, start, end):
    days = pd.date_range(start, end, name="Date")
    table = breakdown(df, "date").drop(columns="Share").set_index("Date")
    return table.reindex(days, fill_value=0).reset_index()


def cell_value(value):
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item"):
        return value.item()
    return value


def style_header(ws, row, count):
    for col in range(1, count + 1):
        cell = ws.cell(row, col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def write_table(ws, table, formats, totals=None, row=1, filters=False, freeze=True):
    for col, name in enumerate(table.columns, 1):
        ws.cell(row, col, name)
    style_header(ws, row, len(table.columns))

    for r, values in enumerate(table.itertuples(index=False), row + 1):
        for c, (name, value) in enumerate(zip(table.columns, values), 1):
            cell = ws.cell(r, c, cell_value(value))
            if name in formats:
                cell.number_format = formats[name]

    if totals:
        r = row + len(table) + 1
        for c, name in enumerate(table.columns, 1):
            cell = ws.cell(r, c, cell_value(totals.get(name)))
            cell.font = BOLD
            cell.border = TOP_LINE
            if name in formats:
                cell.number_format = formats[name]

    if freeze:
        ws.freeze_panes = ws.cell(row + 1, 1)
    if filters:
        ws.auto_filter.ref = f"A{row}:{get_column_letter(len(table.columns))}{row + len(table)}"


def fit_columns(ws, min_width=8, max_width=50):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v is None:
                continue
            if isinstance(v, datetime):
                width = 12
            elif isinstance(v, (int, float)):
                width = len(f"{v:,.2f}") + 3
            else:
                width = len(str(v))
            if cell.row == 1 and ws.auto_filter.ref:
                width += 3  # room for the filter arrow
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), width)
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = min(max(width + 2, min_width), max_width)


def finish_chart(chart, title, axis_format):
    chart.title = title
    chart.title.overlay = False
    chart.legend = None
    chart.varyColors = False
    chart.y_axis.number_format = axis_format
    # openpyxl 3.1 hides both axes unless told otherwise
    chart.x_axis.delete = False
    chart.y_axis.delete = False


def bar_chart(ws, title, value_col, rows, anchor, axis_format, horizontal=False, label_col=1):
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.height, chart.width = 8, 16
    chart.add_data(Reference(ws, min_col=value_col, min_row=1, max_row=rows + 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=label_col, min_row=2, max_row=rows + 1))
    chart.series[0].graphicalProperties.solidFill = CHART_COLOR
    finish_chart(chart, title, axis_format)
    if horizontal:
        # bar charts draw the first row at the bottom, so flip them to put the biggest on top
        # and keep the value axis along the bottom edge
        chart.x_axis.scaling.orientation = "maxMin"
        chart.y_axis.crosses = "max"
    ws.add_chart(chart, anchor)


def daily_chart(ws, rows, anchor, axis_format):
    chart = LineChart()
    chart.height, chart.width = 9, 24
    chart.add_data(Reference(ws, min_col=4, min_row=1, max_row=rows + 1), titles_from_data=True)
    chart.y_axis.crossAx = 500
    chart.x_axis = DateAxis(crossAx=100)
    chart.x_axis.number_format = "d-mmm"
    chart.x_axis.majorTimeUnit = "days"
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=rows + 1))
    line = chart.series[0]
    line.smooth = False
    line.graphicalProperties.line.solidFill = CHART_COLOR
    line.graphicalProperties.line.width = 28000
    finish_chart(chart, "Revenue per day", axis_format)
    ws.add_chart(chart, anchor)


def write_summary(ws, df, issues, stats, money):
    k = kpis(df)
    ws["A1"] = "Sales report"
    ws["A1"].font = Font(bold=True, size=16)

    rows = [
        ("Period", stats["period"], None, None),
        ("Revenue", k["revenue"], money, None),
        ("Orders", k["orders"], INT_FORMAT, None),
        ("Average order value", k["avg_order"], money, None),
        ("Units sold", k["units"], INT_FORMAT, None),
        ("Top product", k["top_product"][0], None, k["top_product"][1]),
        ("Top branch", k["top_branch"][0], None, k["top_branch"][1]),
        ("Generated", datetime.now().replace(microsecond=0), "dd-mmm-yyyy hh:mm", None),
    ]
    for r, (label, value, fmt, revenue) in enumerate(rows, 3):
        ws.cell(r, 1, label).font = BOLD
        cell = ws.cell(r, 2, cell_value(value))
        if fmt:
            cell.number_format = fmt
        if revenue is not None:
            ws.cell(r, 3, cell_value(revenue)).number_format = money

    r += 2
    ws.cell(r, 1, "Data quality").font = SECTION_FONT
    quality = [
        ("Files read", len(stats["files"])),
        ("Rows read", stats["rows_read"]),
        ("Rows dropped", stats["dropped"]),
        ("Rows fixed", stats["fixed"]),
        ("Rows outside period", stats["outside_period"]),
        ("Rows in report", len(df)),
    ]
    for r, (label, value) in enumerate(quality, r + 1):
        ws.cell(r, 1, label).font = BOLD
        ws.cell(r, 2, value).number_format = INT_FORMAT

    r += 2
    ws.cell(r, 1, "What was cleaned").font = SECTION_FONT
    counts = issues.groupby(["issue", "action"], sort=False).size().reset_index(name="Rows")
    counts = counts.rename(columns={"issue": "Issue", "action": "Action"}).sort_values("Rows", ascending=False)
    if counts.empty:
        ws.cell(r + 1, 1, "Nothing to clean")
        r += 2
    else:
        write_table(ws, counts, {"Rows": INT_FORMAT}, row=r + 1, freeze=False)
        r += len(counts) + 3

    ws.cell(r, 1, "Source files").font = SECTION_FONT
    in_report = df["source_file"].value_counts()
    files = pd.DataFrame({
        "File": list(stats["files"]),
        "Rows read": list(stats["files"].values()),
        "Rows in report": [int(in_report.get(name, 0)) for name in stats["files"]],
    })
    write_table(ws, files, {"Rows read": INT_FORMAT, "Rows in report": INT_FORMAT}, row=r + 1, freeze=False)


def write_report(path, df, issues, stats, currency="₹"):
    money = money_format(currency)
    axis = money_format(currency, decimals=0)
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    write_summary(summary, df, issues, stats, money)

    k = kpis(df)
    totals = {"Orders": k["orders"], "Units": k["units"], "Revenue": k["revenue"], "Share": 1}
    formats = {"Orders": INT_FORMAT, "Units": INT_FORMAT, "Revenue": money, "Share": PCT_FORMAT,
               "Avg Order Value": money}

    ws = wb.create_sheet("By Branch")
    table = breakdown(df, "branch")
    table["Avg Order Value"] = table["Revenue"] / table["Orders"]
    write_table(ws, table, formats, totals={"Branch": "Total", **totals, "Avg Order Value": k["avg_order"]})
    bar_chart(ws, "Revenue by branch", 4, len(table), "I2", axis)

    ws = wb.create_sheet("By Category")
    table = breakdown(df, "category")
    write_table(ws, table, formats, totals={"Category": "Total", **totals})
    bar_chart(ws, "Revenue by category", 4, len(table), "H2", axis, horizontal=True)

    ws = wb.create_sheet("Daily")
    table = daily(df, stats["start"], stats["end"])
    write_table(ws, table, {**formats, "Date": DATE_FORMAT},
                totals={"Date": "Total", "Orders": k["orders"], "Units": k["units"], "Revenue": k["revenue"]})
    daily_chart(ws, len(table), "G2", axis)

    ws = wb.create_sheet("Top Products")
    table = breakdown(df, ["product", "category"]).head(10)
    table.insert(0, "Rank", range(1, len(table) + 1))
    write_table(ws, table, formats)
    bar_chart(ws, "Top 10 products by revenue", 6, len(table), "I2", axis, horizontal=True, label_col=2)

    ws = wb.create_sheet("Clean Data")
    data = df.rename(columns=heading)
    write_table(ws, data, {"Date": DATE_FORMAT, "Unit Price": money, "Total": money}, filters=True)

    ws = wb.create_sheet("Data Issues")
    table = issues.rename(columns=heading)
    write_table(ws, table, {"Source Row": "0"}, filters=not issues.empty)
    if issues.empty:
        ws.cell(2, 1, "No problems found")

    for ws in wb.worksheets:
        fit_columns(ws)
        # openpyxl turns any string starting with "=" into a live formula. The report has no
        # formulas of its own, so anything like that came from the input files and stays text.
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
    wb.save(path)

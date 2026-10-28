from pathlib import Path

import pytest
from openpyxl import load_workbook

import build_report
import mailer
from build_report import build, main
from report import SHEETS, write_report

SAMPLES = Path(__file__).resolve().parent.parent / "data" / "raw"


def labels(ws):
    return {row[0].value: row[1].value for row in ws.iter_rows(max_col=2) if row[0].value}


def total_row(ws):
    last = [c.value for c in ws[ws.max_row]]
    assert last[0] == "Total"
    return dict(zip([c.value for c in ws[1]], last))


@pytest.fixture(scope="module")
def september(tmp_path_factory):
    data, issues, stats = build(SAMPLES, month="2026-09")
    path = tmp_path_factory.mktemp("report") / "report.xlsx"
    write_report(path, data, issues, stats)
    return data, issues, stats, load_workbook(path)


def test_all_sheets_present(september):
    *_, wb = september
    assert wb.sheetnames == SHEETS


def test_charts_present(september):
    *_, wb = september
    for name in ["By Branch", "By Category", "Daily", "Top Products"]:
        assert len(wb[name]._charts) == 1, name


def test_totals_match_clean_data(september):
    data, _, _, wb = september
    revenue = round(data["total"].sum(), 2)
    orders = len(data[["branch", "order_id"]].drop_duplicates())

    summary = labels(wb["Summary"])
    assert summary["Revenue"] == pytest.approx(revenue)
    assert summary["Orders"] == orders
    assert summary["Average order value"] == pytest.approx(revenue / orders, abs=0.01)
    assert summary["Units sold"] == data["quantity"].sum()

    for name in ["By Branch", "By Category", "Daily"]:
        assert total_row(wb[name])["Revenue"] == pytest.approx(revenue), name
        assert total_row(wb[name])["Orders"] == orders, name

    branch_rows = list(wb["By Branch"].iter_rows(min_row=2, max_row=wb["By Branch"].max_row - 1, values_only=True))
    assert sum(r[3] for r in branch_rows) == pytest.approx(revenue)

    daily = list(wb["Daily"].iter_rows(min_row=2, max_row=wb["Daily"].max_row - 1, values_only=True))
    assert len(daily) == 30
    assert sum(r[3] for r in daily) == pytest.approx(revenue)

    clean_rows = list(wb["Clean Data"].iter_rows(min_row=2, values_only=True))
    assert len(clean_rows) == len(data)
    assert sum(r[7] for r in clean_rows) == pytest.approx(revenue)


def test_every_row_is_accounted_for(september):
    data, issues, stats, wb = september
    dropped = (issues["action"] == "dropped").sum()
    assert stats["rows_read"] == len(data) + dropped + stats["outside_period"]

    issue_rows = wb["Data Issues"].max_row - 1
    assert issue_rows == len(issues)
    assert labels(wb["Summary"])["Rows dropped"] == dropped


def test_sample_duplicates_are_removed():
    data, issues, _ = build(SAMPLES)
    counts = issues.groupby("issue").size()
    assert counts["blank row"] == 15
    assert counts["exact duplicate"] == 19
    assert counts["repeated order line"] == 12
    assert not data.duplicated(["branch", "order_id", "product"]).any()


def test_month_filter():
    everything, _, _ = build(SAMPLES)
    data, _, stats = build(SAMPLES, month="2026-09")

    assert stats["period"] == "September 2026"
    assert (data["date"].dt.month == 9).all()
    assert stats["outside_period"] == (everything["date"].dt.month == 8).sum()
    assert stats["outside_period"] > 0
    assert len(data) + stats["outside_period"] == len(everything)


def test_cli_writes_report(tmp_path):
    out = tmp_path / "out" / "report.xlsx"
    assert main([str(SAMPLES), "-o", str(out), "--month", "2026-09", "--currency", "Rs "]) == 0
    ws = load_workbook(out)["Summary"]
    revenue_cell = next(row[1] for row in ws.iter_rows(max_col=2) if row[0].value == "Revenue")
    assert revenue_cell.number_format == '"Rs "#,##0.00'


def test_cli_rejects_empty_month(tmp_path):
    assert main([str(SAMPLES), "-o", str(tmp_path / "r.xlsx"), "--month", "2025-01"]) == 1


class FakeSMTP:
    sent = []

    def __init__(self, host, port, timeout):
        pass

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def starttls(self):
        pass

    def login(self, user, password):
        pass

    def send_message(self, msg):
        self.sent.append(msg)


@pytest.fixture
def smtp(monkeypatch):
    FakeSMTP.sent = []
    monkeypatch.setattr(mailer.smtplib, "SMTP", FakeSMTP)
    monkeypatch.setattr(build_report, "load_dotenv", lambda: None)
    for key in ["SMTP_USER", "SMTP_PASSWORD", "SMTP_FROM"]:
        monkeypatch.delenv(key, raising=False)
    return FakeSMTP.sent


def test_email_skipped_when_not_configured(tmp_path, smtp):
    out = tmp_path / "report.xlsx"
    assert main([str(SAMPLES), "-o", str(out), "--email", "owner@example.com"]) == 0
    assert out.exists()
    assert smtp == []


def test_email_sends_report(tmp_path, smtp, monkeypatch):
    monkeypatch.setenv("SMTP_USER", "reports@example.com")
    monkeypatch.setenv("SMTP_PASSWORD", "test-password")
    out = tmp_path / "report.xlsx"

    assert main([str(SAMPLES), "-o", str(out), "--month", "2026-09", "--email", "owner@example.com"]) == 0
    msg = smtp[0]
    assert msg["To"] == "owner@example.com"
    assert msg["Subject"] == "Sales report: September 2026"
    attachment = next(msg.iter_attachments())
    assert attachment.get_filename() == "report.xlsx"
    assert attachment.get_content() == out.read_bytes()


def test_formula_text_in_input_stays_text(tmp_path):
    data, issues, stats = build(SAMPLES, month="2026-09")
    data = data.copy()
    data.iloc[0, data.columns.get_loc("product")] = "=1+1"
    path = tmp_path / "report.xlsx"
    write_report(path, data, issues, stats)

    ws = load_workbook(path)["Clean Data"]
    cells = [c for row in ws.iter_rows() for c in row if c.value == "=1+1"]
    assert cells and all(c.data_type == "s" for c in cells)

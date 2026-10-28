"""Writes three messy sales exports for a fictional electronics accessories chain.

Each branch uses different billing software, so every file has its own column
names, date format and way of writing money. The random generator is seeded,
so the files come out the same on every run.

    python samples/make_messy_data.py [output folder]
"""
import csv
import random
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

OUT_DIR = Path(__file__).resolve().parent.parent / "data" / "raw"

# name, category, price in rupees, how often it sells
PRODUCTS = [
    ("USB-C Cable 1m", "Mobile Accessories", 299, 10),
    ("Fast Charger 20W", "Mobile Accessories", 899, 7),
    ("Phone Case", "Mobile Accessories", 349, 9),
    ("Tempered Glass", "Mobile Accessories", 199, 10),
    ("Power Bank 10000mAh", "Mobile Accessories", 1299, 5),
    ("Wireless Earbuds", "Audio", 1999, 6),
    ("Neckband Earphones", "Audio", 1099, 5),
    ("Bluetooth Speaker", "Audio", 2499, 3),
    ("Smart Watch", "Wearables", 3499, 3),
    ("Fitness Band", "Wearables", 1799, 2),
    ("Wireless Mouse", "Computer Accessories", 599, 5),
    ("Keyboard Mouse Combo", "Computer Accessories", 1149, 3),
    ("Pendrive 64GB", "Computer Accessories", 549, 6),
    ("HDMI Cable 1.5m", "Computer Accessories", 399, 3),
    ("Laptop Backpack", "Computer Accessories", 1499, 2),
    ("LED Desk Lamp", "Home & Office", 1199, 2),
    ("Extension Board 4 Socket", "Home & Office", 649, 4),
    ("Wi-Fi Router", "Home & Office", 2299, 2),
]


def make_sales(rng, branch, busy, days=None):
    days = days or [date(2026, 9, d) for d in range(1, 31)]
    weights = [p[3] for p in PRODUCTS]
    sales, order_no = [], 0
    for day in days:
        weekend = 1.4 if day.weekday() >= 5 else 1
        for _ in range(round(rng.randint(5, 9) * busy * weekend)):
            order_no += 1
            basket = {}
            for _ in range(rng.choice([1, 1, 1, 2, 2, 3])):
                product = rng.choices(PRODUCTS, weights)[0]
                basket[product[0]] = product
            for name, category, price, _ in basket.values():
                sales.append({"date": day, "order_no": order_no, "branch": branch, "product": name,
                              "category": category, "qty": rng.choice([1, 1, 1, 1, 2, 2, 3]), "price": price})
    return sales


def scruffy(rng, text, chance=0.08):
    if rng.random() > chance:
        return text
    return rng.choice([text.upper(), text.lower(), f" {text}", f"{text}  ", text.replace(" ", "  ", 1)])


def add_duplicates(rng, rows, count):
    for _ in range(count):
        i = rng.randrange(len(rows))
        rows.insert(min(i + rng.randint(1, 20), len(rows)), list(rows[i]))


def add_blanks(rng, rows, count, width):
    for _ in range(count):
        rows.insert(rng.randrange(len(rows)), [""] * width if rng.random() < 0.5 else [])


def write_csv(path, header, rows, encoding):
    with open(path, "w", newline="", encoding=encoding) as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def pune(rng, path):
    # POS export: dd/mm/yyyy dates, no amount column, saved from Excel (with BOM)
    rows = []
    for s in make_sales(rng, "Pune", 1.3):
        rate = rng.choice([f"₹{s['price']:,.2f}", f"{s['price']:.2f}", f"₹ {s['price']}"])
        rows.append([f"PUN-{10000 + s['order_no']}", s["date"].strftime("%d/%m/%Y"), scruffy(rng, "Pune", 0.1),
                     scruffy(rng, s["product"]), scruffy(rng, s["category"], 0.05), str(s["qty"]), rate])

    picks = rng.sample(range(len(rows)), 11)
    for i in picks[:4]:
        rows[i][5] = ""
    for i in picks[4:7]:
        rows[i][4] = ""
    for i in picks[7:9]:
        rows[i][5] += " pcs"
    for i in picks[9:]:
        rows[i][1] = "31/09/2026"

    add_duplicates(rng, rows, 8)
    add_blanks(rng, rows, 6, 7)
    write_csv(path, ["Bill No", "Bill Date", "Outlet", "Item Name", "Category", "Qty", "Rate"], rows, "utf-8-sig")
    return len(rows)


def rupees(rng, value):
    return rng.choice([f"Rs {value:g}", f"Rs. {value:,.2f}", f"{value:g}/-", f"{value:.2f}"])


def nashik_date(rng, day):
    return day.isoformat() if rng.random() < 0.7 else f"{day.day} {day:%b %Y}"


def nashik_row(rng, s):
    return [f"NSK/26-27/{450 + s['order_no']:04d}", nashik_date(rng, s["date"]),
            scruffy(rng, "Nashik", 0.1), scruffy(rng, s["product"]), str(s["qty"]),
            rupees(rng, s["price"]), rupees(rng, s["price"] * s["qty"]), "18"]


def nashik(rng, path):
    # accounting software export: ISO and "3 Sep 2026" dates mixed, "Rs" amounts, no category
    sales = make_sales(rng, "Nashik", 0.9)
    rows = [nashik_row(rng, s) for s in sales]

    picks = rng.sample(range(len(rows)), 6)
    for i in picks[:5]:
        rows[i][4] = ""
    rows[picks[5]][1] = "2026-09-31"

    # someone exported a few days twice and pasted the second copy underneath,
    # so the same invoice lines come back with different formatting
    start = rng.randrange(len(sales) - 12)
    for s in sales[start:start + 12]:
        again = nashik_row(rng, s)
        again[0] = again[0].lower()
        rows.append(again)

    add_duplicates(rng, rows, 5)
    add_blanks(rng, rows, 4, 8)
    header = ["Invoice Number", "Invoice Date", "Branch Name", "Product", "Quantity", "Unit Price",
              "Net Amount", "GST %"]
    write_csv(path, header, rows, "utf-8")
    return len(rows)


def nagpur(rng, path):
    # Excel register with a title above the header, real date cells mixed with text dates
    late = make_sales(rng, "Nagpur", 0.15, days=[date(2026, 8, 29), date(2026, 8, 31)])
    sales = make_sales(rng, "Nagpur", 1.1)
    rows = []
    for n, s in enumerate(late + sales):
        order_id = 50000 + s["order_no"] if n >= len(late) else 49900 + s["order_no"]
        when = datetime.combine(s["date"], datetime.min.time())
        if rng.random() < 0.15:
            when = rng.choice([s["date"].strftime("%d-%m-%Y"), f"{s['date'].day} {s['date']:%b %Y}"])
        price = s["price"] if rng.random() > 0.1 else f"₹ {s['price']:,.2f}"
        rows.append([when, order_id, scruffy(rng, "Nagpur", 0.1), scruffy(rng, s["product"]),
                     scruffy(rng, s["category"], 0.05), s["qty"], price, s["price"] * s["qty"]])

    picks = rng.sample(range(len(late), len(rows)), 9)
    for i in picks[:2]:
        rows[i][1] = None
    for i in picks[2:5]:
        rows[i][5] = None
    for i in picks[5:8]:
        rows[i][7] = None
    rows[picks[8]][5] = 0

    add_duplicates(rng, rows, 6)
    add_blanks(rng, rows, 5, 8)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Register"
    ws.append(["Sales register: Nagpur branch, September 2026"])
    ws.append([])
    ws.append(["Date", "Order ID", "Store", "Product Description", "Product Category", "Units", "Price",
               "Total Amount"])
    for row in rows:
        ws.append([None if v == "" else v for v in row])
        if row and isinstance(row[0], datetime):
            ws.cell(ws.max_row, 1).number_format = "dd-mm-yyyy"
    wb.save(path)
    return len(rows)


def main(out_dir=OUT_DIR):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    rng = random.Random(2026)
    for name, func in [("pune_pos_export.csv", pune), ("nashik_billing.csv", nashik), ("nagpur_sales.xlsx", nagpur)]:
        count = func(rng, out_dir / name)
        print(f"{out_dir / name}: {count} rows")


if __name__ == "__main__":
    main(*sys.argv[1:])
